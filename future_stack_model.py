#!/usr/bin/env python3
"""
Future-Stack Trading Model v2.7 — scoring engine.

Single source of truth for the 10-category, 100-point conviction rubric
used all session, replacing the one-off inline Python snippets (rebuilt by
hand for every new stock or screen) with one reusable, tested tool.

GOLDEN RULE: a score measures conviction/quality in a business. It is
NEVER a buy or sell signal, and this tool does not recommend trades — it
only computes and records scores.

Categories (fixed order, caps sum to exactly 100):
  1. Structural Sector Future          /10
  2. Future-Stack Role                  /7
  3. Durable Moat                      /14
  4. Revenue Quality & Growth          /12
  5. Margin Path / Unit Economics      /10
  6. FCF & Balance Sheet               /12
  7. Management & Capital Allocation    /8
  8. Proof / Milestone Quality         /14
  9. Valuation                         /8
 10. Macro / Rate Sensitivity          /5

Grade bands: 92-100 A+ | 86-91 A | 80-85 A- | 74-79 B+ | 68-73 B |
             60-67 C+ | 50-59 C | <50 Reject

Usage:
  python3 future_stack_model.py score --ticker AVPT --name "AvePoint" \\
      --scores 8,5,8,10,8,11,7,8,5,3 --notes "Independent scorecard."

  python3 future_stack_model.py batch candidates.csv
      (CSV columns: ticker,name,scores,notes -- scores is a quoted,
       comma-separated list of 10 numbers in the fixed category order)

  python3 future_stack_model.py append --tracker trading_model_tracker.xlsx \\
      --ticker AVPT --name "AvePoint" --scores 8,5,8,10,8,11,7,8,5,3 \\
      --scope "Current portfolio holding" --notes "..."
      (inserts a new row into the tracker's "Per-Stock Scorecards" tab,
       ABOVE the TODO row, with live formulas -- run the xlsx skill's
       recalc.py afterward so the Total/Grade formulas get cached values)

  python3 future_stack_model.py --show-categories
"""
import argparse
import csv
import sys

CATEGORIES = [
    ("Structural Sector Future", 10),
    ("Future-Stack Role", 7),
    ("Durable Moat", 14),
    ("Revenue Quality & Growth", 12),
    ("Margin Path / Unit Economics", 10),
    ("FCF & Balance Sheet", 12),
    ("Management & Capital Allocation", 8),
    ("Proof / Milestone Quality", 14),
    ("Valuation", 8),
    ("Macro / Rate Sensitivity", 5),
]
assert sum(cap for _, cap in CATEGORIES) == 100, "Category caps must sum to exactly 100."

# (minimum score inclusive, grade) -- checked highest-first
GRADE_BANDS = [
    (92, "A+"), (86, "A"), (80, "A-"), (74, "B+"),
    (68, "B"), (60, "C+"), (50, "C"), (0, "Reject"),
]


class ScoreError(ValueError):
    pass


def grade(total):
    for min_score, g in GRADE_BANDS:
        if total >= min_score:
            return g
    return "Reject"  # unreachable (0 is the floor) but keeps the function total


def validate_scores(scores):
    if len(scores) != len(CATEGORIES):
        raise ScoreError(f"Expected {len(CATEGORIES)} category scores, got {len(scores)}: {scores}")
    for (name, cap), s in zip(CATEGORIES, scores):
        if not (0 <= s <= cap):
            raise ScoreError(f"'{name}' scored {s}, out of its allowed range 0-{cap}")


def score_stock(ticker, scores, name="", notes=""):
    """Validate and score one stock. scores: list/tuple of 10 numbers in CATEGORIES order."""
    scores = [float(s) if float(s) != int(s) else int(s) for s in scores]
    validate_scores(scores)
    total = sum(scores)
    return {
        "ticker": ticker.strip().upper(),
        "name": name.strip(),
        "scores": dict(zip((c[0] for c in CATEGORIES), scores)),
        "score_list": scores,
        "total": total,
        "grade": grade(total),
        "notes": notes.strip(),
    }


def format_scorecard(result):
    header = f"{result['ticker']} — {result['name']}" if result["name"] else result["ticker"]
    lines = [header]
    for (cname, cap), s in zip(CATEGORIES, result["score_list"]):
        lines.append(f"  {cname:<36} {s:>5}/{cap}")
    lines.append(f"  {'-' * 44}")
    lines.append(f"  {'TOTAL':<36} {result['total']:>5}/100   Grade: {result['grade']}")
    if result["notes"]:
        lines.append(f"  Notes: {result['notes']}")
    return "\n".join(lines)


def parse_scores_arg(s):
    try:
        return [float(x) for x in s.split(",")]
    except ValueError:
        raise ScoreError(f"Could not parse --scores '{s}' as 10 comma-separated numbers.")


# ---------------------------------------------------------------------------
# xlsx tracker integration
# ---------------------------------------------------------------------------

def append_to_tracker(tracker_path, result, scope=""):
    """
    Insert `result` as a new row in the tracker's "Per-Stock Scorecards"
    sheet, directly above the first row whose Ticker column reads "TODO"
    (if present), otherwise at the first fully empty row. Writes a live
    SUM formula for Total and an INDEX/MATCH formula for Grade, matching
    the sheet's existing convention. Caller should run the xlsx skill's
    recalc.py afterward (openpyxl can't compute formula results itself).
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(tracker_path)
    if "Per-Stock Scorecards" not in wb.sheetnames:
        raise RuntimeError('Tracker is missing a "Per-Stock Scorecards" sheet -- wrong file?')
    ws = wb["Per-Stock Scorecards"]

    header_row = 1
    todo_row = None
    insert_row = None
    for r in range(2, ws.max_row + 2):
        val = ws.cell(row=r, column=1).value
        if val == "TODO":
            todo_row = r
            break
        if val is None and ws.cell(row=r, column=2).value is None:
            insert_row = r
            break
    target_row = todo_row if todo_row is not None else (insert_row or ws.max_row + 1)

    if todo_row is not None:
        # openpyxl's insert_rows does not reliably re-anchor merged cell
        # ranges that start at or below the insertion point -- left alone,
        # the TODO row's merge (columns B:Q) silently eats every value
        # written into the new row above it on save, with no error raised.
        # Unmerge first, insert, then re-merge one row lower.
        ranges_to_shift = [str(mc) for mc in list(ws.merged_cells.ranges) if mc.min_row >= target_row]
        for rng in ranges_to_shift:
            ws.unmerge_cells(rng)
        ws.insert_rows(target_row)
        from openpyxl.utils.cell import range_boundaries
        for rng in ranges_to_shift:
            min_col, min_row, max_col, max_row = range_boundaries(rng)
            ws.merge_cells(start_row=min_row + 1, start_column=min_col,
                            end_row=max_row + 1, end_column=max_col)

    first_cat_col = 5  # column E: Ticker(A) Company(B) Date(C) Scope(D) then categories start
    last_cat_col = first_cat_col + len(CATEGORIES) - 1
    total_col = last_cat_col + 1
    grade_col = total_col + 1
    notes_col = grade_col + 1

    ws.cell(row=target_row, column=1, value=result["ticker"])
    ws.cell(row=target_row, column=2, value=result["name"])
    ws.cell(row=target_row, column=3, value="(auto-appended by future_stack_model.py)")
    ws.cell(row=target_row, column=4, value=scope)
    for i, s in enumerate(result["score_list"]):
        ws.cell(row=target_row, column=first_cat_col + i, value=s)

    first_letter = get_column_letter(first_cat_col)
    last_letter = get_column_letter(last_cat_col)
    ws.cell(row=target_row, column=total_col,
            value=f"=SUM({first_letter}{target_row}:{last_letter}{target_row})")
    ws.cell(row=target_row, column=grade_col,
            value=f"=INDEX(Reference!$C$2:$C$9,MATCH({get_column_letter(total_col)}{target_row},"
                  f"Reference!$B$2:$B$9,1))")
    ws.cell(row=target_row, column=notes_col, value=result["notes"])

    wb.save(tracker_path)
    return target_row


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--show-categories", action="store_true", help="Print the category list and caps, then exit.")
    sub = p.add_subparsers(dest="command")

    sc = sub.add_parser("score", help="Score one stock and print the scorecard.")
    sc.add_argument("--ticker", required=True)
    sc.add_argument("--name", default="")
    sc.add_argument("--scores", required=True, help="10 comma-separated numbers, in the fixed category order.")
    sc.add_argument("--notes", default="")

    bt = sub.add_parser("batch", help="Score multiple stocks from a CSV and print them sorted by total, highest first.")
    bt.add_argument("csv_path")
    bt.add_argument("--out", help="Optional path to write a results CSV (ticker,name,total,grade,notes).")

    ap = sub.add_parser("append", help="Score one stock and insert it into a tracker workbook's Per-Stock Scorecards tab.")
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--ticker", required=True)
    ap.add_argument("--name", default="")
    ap.add_argument("--scores", required=True)
    ap.add_argument("--scope", default="")
    ap.add_argument("--notes", default="")

    args = p.parse_args()

    if args.show_categories:
        for name, cap in CATEGORIES:
            print(f"  {name:<36} /{cap}")
        print(f"  {'TOTAL CAP':<36} /100")
        return

    if args.command == "score":
        try:
            result = score_stock(args.ticker, parse_scores_arg(args.scores), args.name, args.notes)
        except ScoreError as e:
            sys.exit(f"Score error: {e}")
        print(format_scorecard(result))

    elif args.command == "batch":
        results = []
        with open(args.csv_path, newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                try:
                    scores = parse_scores_arg(row["scores"])
                    results.append(score_stock(row["ticker"], scores, row.get("name", ""), row.get("notes", "")))
                except ScoreError as e:
                    print(f"SKIPPED {row.get('ticker', '?')}: {e}", file=sys.stderr)
        results.sort(key=lambda r: r["total"], reverse=True)
        for r in results:
            print(f"{r['ticker']:<8} {r['total']:>4}/100  {r['grade']:<7} {r['name']}")
        if args.out:
            with open(args.out, "w", newline="") as f:
                w = csv.writer(f)
                w.writerow(["ticker", "name", "total", "grade", "notes"])
                for r in results:
                    w.writerow([r["ticker"], r["name"], r["total"], r["grade"], r["notes"]])
            print(f"\nWrote {len(results)} rows to {args.out}")

    elif args.command == "append":
        try:
            result = score_stock(args.ticker, parse_scores_arg(args.scores), args.name, args.notes)
        except ScoreError as e:
            sys.exit(f"Score error: {e}")
        row = append_to_tracker(args.tracker, result, scope=args.scope)
        print(format_scorecard(result))
        print(f"\nInserted at row {row} in {args.tracker}'s 'Per-Stock Scorecards' tab.")
        print("Run the xlsx skill's recalc.py on the tracker next so Total/Grade get cached values.")

    else:
        p.print_help()

    print("\nGolden Rule: a score measures conviction/quality. It is never a buy or sell signal.")


if __name__ == "__main__":
    main()
