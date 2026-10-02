#!/usr/bin/env python3
"""
AI Portfolio Management System — the portfolio-level engine, as opposed to
future_stack_model.py's per-stock engine. Same reporting/flagging-only
discipline: this never recommends or gates a trade, it only scores
portfolio construction, analyzes concentration/diversification, and tracks
open red flags.

GOLDEN RULE (same as the stock engine): a score measures structural
quality/risk. It is never a buy or sell signal.

Covers three things:
  1. The Portfolio Construction Score -- a separate 7-category, 100-point
     rubric scoring how the WHOLE portfolio is built (diversification,
     concentration, quality mix, liquidity, leverage, open flags, thematic
     durability) -- distinct from any individual stock's Future-Stack score.
  2. A holdings-based concentration/diversification analyzer -- computable
     from real position data (HHI, top-position weight, theme spread),
     offered as a data-derived SUGGESTION for two of the rubric's
     categories, not a binding answer.
  3. Red-flag log management -- add, list, and resolve open flags, reading
     and writing directly to a tracker workbook's "Red Flag Log" sheet so
     there's one durable record instead of flags living only in chat memory.

Usage:
  python3 portfolio_management_system.py score --scores 10,10,11,3,6,8,12 --notes "..."

  python3 portfolio_management_system.py analyze-holdings holdings.csv
      (CSV columns: ticker,value,theme)

  python3 portfolio_management_system.py flags list --tracker trading_model_tracker.xlsx
  python3 portfolio_management_system.py flags add --tracker trading_model_tracker.xlsx \\
      --flag "New SEC inquiry" --ticker XYZ --status "Open / Monitoring" \\
      --notes "..." --source "..."
  python3 portfolio_management_system.py flags resolve --tracker trading_model_tracker.xlsx \\
      --ticker WLDN --match "dilution" --status "Resolved" --notes "..."

  python3 portfolio_management_system.py append --tracker trading_model_tracker.xlsx \\
      --scores 10,10,11,3,6,8,12 --notes "..."
      (logs a new dated row to a running Score History table in the
       tracker's "Portfolio Construction Score" sheet, creating that table
       if it doesn't exist yet)
"""
import argparse
import csv
import datetime
import math
import sys

CATEGORIES = [
    ("Diversification across uncorrelated themes", 20),
    ("Single-position concentration risk", 15),
    ("Quality/stability mix (barbell construction)", 15),
    ("Liquidity / cash buffer", 10),
    ("Leverage / instrument-risk exposure", 10),
    ("Known open red flags on file", 15),
    ("Thematic conviction quality (durability of megatrends)", 15),
]
assert sum(cap for _, cap in CATEGORIES) == 100, "Category caps must sum to exactly 100."

# Reusing the same letter-grade bands as the per-stock Future-Stack model for
# a consistent vocabulary across both engines (both are /100 scores) -- this
# is a borrowed label for convenience, NOT a claim that a portfolio-construction
# 75 and a stock-conviction 75 mean the same thing. Say so whenever this
# matters.
GRADE_BANDS = [
    (92, "A+"), (86, "A"), (80, "A-"), (74, "B+"),
    (68, "B"), (60, "C+"), (50, "C"), (0, "Reject"),
]


class ScoreError(ValueError):
    pass


def grade(total):
    for min_score, g in GRADE_BANDS:
        if total >= min_score:
            return g
    return "Reject"


def validate_scores(scores):
    if len(scores) != len(CATEGORIES):
        raise ScoreError(f"Expected {len(CATEGORIES)} category scores, got {len(scores)}: {scores}")
    for (name, cap), s in zip(CATEGORIES, scores):
        if not (0 <= s <= cap):
            raise ScoreError(f"'{name}' scored {s}, out of its allowed range 0-{cap}")


def score_portfolio(scores, notes=""):
    scores = [float(s) if float(s) != int(s) else int(s) for s in scores]
    validate_scores(scores)
    total = sum(scores)
    return {
        "scores": dict(zip((c[0] for c in CATEGORIES), scores)),
        "score_list": scores,
        "total": total,
        "grade": grade(total),
        "notes": notes.strip(),
    }


def format_scorecard(result):
    lines = ["Portfolio Construction Score"]
    for (cname, cap), s in zip(CATEGORIES, result["score_list"]):
        lines.append(f"  {cname:<55} {s:>5}/{cap}")
    lines.append(f"  {'-' * 63}")
    lines.append(f"  {'TOTAL':<55} {result['total']:>5}/100   Grade: {result['grade']}")
    if result["notes"]:
        lines.append(f"  Notes: {result['notes']}")
    return "\n".join(lines)


def parse_scores_arg(s):
    try:
        return [float(x) for x in s.split(",")]
    except ValueError:
        raise ScoreError(f"Could not parse --scores '{s}' as {len(CATEGORIES)} comma-separated numbers.")


# ---------------------------------------------------------------------------
# Holdings-based concentration / diversification analysis
# ---------------------------------------------------------------------------

def analyze_holdings(rows):
    """
    rows: list of dicts with 'ticker', 'value' (float), 'theme' (str).
    Returns HHI, top-position weight, per-theme weights, and a SUGGESTED
    (not binding) score for the two rubric categories this can actually be
    computed from: Single-position concentration and Diversification.
    """
    total_value = sum(r["value"] for r in rows)
    if total_value <= 0:
        raise ValueError("Total portfolio value must be positive.")

    weights = [(r["ticker"], r["value"] / total_value) for r in rows]
    weights.sort(key=lambda x: x[1], reverse=True)
    top_ticker, top_weight = weights[0]

    hhi = sum(w ** 2 for _, w in weights)  # 1/n (perfectly diversified) .. 1 (single position)
    effective_n = 1 / hhi if hhi > 0 else 0  # "effective number of equal-sized positions"

    theme_totals = {}
    for r in rows:
        theme_totals[r["theme"]] = theme_totals.get(r["theme"], 0) + r["value"]
    theme_weights = {t: v / total_value for t, v in theme_totals.items()}
    n_themes = len(theme_weights)
    theme_hhi = sum(w ** 2 for w in theme_weights.values())
    largest_theme, largest_theme_weight = max(theme_weights.items(), key=lambda kv: kv[1])

    # Suggested sub-scores -- simple, transparent, adjustable thresholds.
    # Concentration (cap 15): penalize a large top single position.
    if top_weight <= 0.08:
        suggested_concentration = 15
    elif top_weight <= 0.12:
        suggested_concentration = 12
    elif top_weight <= 0.18:
        suggested_concentration = 9
    elif top_weight <= 0.25:
        suggested_concentration = 6
    else:
        suggested_concentration = 3

    # Diversification (cap 20): penalize a large largest-theme weight / few themes.
    if n_themes >= 6 and largest_theme_weight <= 0.25:
        suggested_diversification = 18
    elif n_themes >= 4 and largest_theme_weight <= 0.35:
        suggested_diversification = 13
    elif n_themes >= 3 and largest_theme_weight <= 0.50:
        suggested_diversification = 9
    elif largest_theme_weight <= 0.65:
        suggested_diversification = 6
    else:
        suggested_diversification = 3

    return {
        "total_value": total_value,
        "n_positions": len(rows),
        "top_ticker": top_ticker, "top_weight": top_weight,
        "hhi": hhi, "effective_n_positions": effective_n,
        "theme_weights": dict(sorted(theme_weights.items(), key=lambda kv: kv[1], reverse=True)),
        "n_themes": n_themes, "theme_hhi": theme_hhi,
        "largest_theme": largest_theme, "largest_theme_weight": largest_theme_weight,
        "suggested_concentration_score": suggested_concentration,
        "suggested_diversification_score": suggested_diversification,
    }


def format_holdings_analysis(a):
    lines = [
        f"Portfolio: {a['n_positions']} positions, total value ${a['total_value']:,.2f}",
        f"Largest single position: {a['top_ticker']} at {a['top_weight']:.1%} of portfolio",
        f"Herfindahl-Hirschman Index (positions): {a['hhi']:.4f} "
        f"(effective ~{a['effective_n_positions']:.1f} equal-sized positions)",
        "",
        f"Themes represented: {a['n_themes']}",
    ]
    for theme, w in a["theme_weights"].items():
        lines.append(f"  {theme:<40} {w:>6.1%}")
    lines.append(f"Largest theme: {a['largest_theme']} at {a['largest_theme_weight']:.1%}")
    lines.append("")
    lines.append(f"SUGGESTED (not binding -- a data-derived starting point, review before using):")
    lines.append(f"  Single-position concentration risk: {a['suggested_concentration_score']}/15")
    lines.append(f"  Diversification across uncorrelated themes: {a['suggested_diversification_score']}/20")
    lines.append("  (\"Theme\" here is whatever label your CSV used -- it only measures spread across")
    lines.append("   the labels you gave it, not whether those themes are ACTUALLY uncorrelated in a")
    lines.append("   selloff. That judgment call is still yours.)")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Red Flag Log (reads/writes the tracker workbook directly)
# ---------------------------------------------------------------------------

FLAG_HEADERS = ["Flag", "Ticker", "Status", "Last Checked", "Finding / Notes", "Source"]


def _open_flag_sheet(tracker_path):
    import openpyxl
    wb = openpyxl.load_workbook(tracker_path)
    if "Red Flag Log" not in wb.sheetnames:
        raise RuntimeError('Tracker is missing a "Red Flag Log" sheet -- wrong file?')
    return wb, wb["Red Flag Log"]


def list_flags(tracker_path, ticker=None, status_contains=None):
    wb, ws = _open_flag_sheet(tracker_path)
    out = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, 7)]
        if row[0] is None and row[1] is None:
            continue
        if ticker and (row[1] or "").strip().upper() != ticker.strip().upper():
            continue
        if status_contains and status_contains.lower() not in (row[2] or "").lower():
            continue
        out.append(dict(zip(FLAG_HEADERS, row)))
    return out


def add_flag(tracker_path, flag, ticker, status, notes="", source="", last_checked=None):
    wb, ws = _open_flag_sheet(tracker_path)
    r = 2
    while ws.cell(row=r, column=1).value is not None or ws.cell(row=r, column=2).value is not None:
        r += 1
    last_checked = last_checked or datetime.date.today().isoformat()
    values = [flag, ticker.strip().upper(), status, last_checked, notes, source]
    for i, v in enumerate(values, start=1):
        ws.cell(row=r, column=i, value=v)
    wb.save(tracker_path)
    return r


def resolve_flag(tracker_path, ticker, match, new_status, notes_append="", last_checked=None):
    """Find row(s) for `ticker` whose Flag/Notes text contains `match` (case-insensitive),
    update Status + Last Checked, and append to the Notes field rather than overwriting it."""
    wb, ws = _open_flag_sheet(tracker_path)
    last_checked = last_checked or datetime.date.today().isoformat()
    updated_rows = []
    for r in range(2, ws.max_row + 1):
        flag_val = ws.cell(row=r, column=1).value or ""
        ticker_val = ws.cell(row=r, column=2).value or ""
        notes_val = ws.cell(row=r, column=5).value or ""
        if ticker_val.strip().upper() != ticker.strip().upper():
            continue
        if match.lower() not in (flag_val + " " + notes_val).lower():
            continue
        ws.cell(row=r, column=3, value=new_status)
        ws.cell(row=r, column=4, value=last_checked)
        if notes_append:
            combined = (notes_val + " | " if notes_val else "") + f"[{last_checked}] {notes_append}"
            ws.cell(row=r, column=5, value=combined)
        updated_rows.append(r)
    if not updated_rows:
        raise RuntimeError(f"No Red Flag Log row found for ticker={ticker!r} matching text {match!r}")
    wb.save(tracker_path)
    return updated_rows


# ---------------------------------------------------------------------------
# Portfolio Construction Score -- dated history log in the tracker
# ---------------------------------------------------------------------------

HISTORY_MARKER = "Score History"


def append_score_history(tracker_path, result, as_of=None):
    import openpyxl
    wb = openpyxl.load_workbook(tracker_path)
    if "Portfolio Construction Score" not in wb.sheetnames:
        raise RuntimeError('Tracker is missing a "Portfolio Construction Score" sheet -- wrong file?')
    ws = wb["Portfolio Construction Score"]
    as_of = as_of or datetime.date.today().isoformat()

    marker_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == HISTORY_MARKER:
            marker_row = r
            break

    if marker_row is None:
        # Create the history table from scratch, a couple of rows below whatever's there.
        start_row = ws.max_row + 3
        ws.cell(row=start_row, column=1, value=HISTORY_MARKER)
        header_row = start_row + 1
        headers = ["Date"] + [c[0] for c in CATEGORIES] + ["Total", "Grade", "Notes"]
        for i, h in enumerate(headers, start=1):
            ws.cell(row=header_row, column=i, value=h)
        data_row = header_row + 1
    else:
        header_row = marker_row + 1
        data_row = header_row + 1
        while ws.cell(row=data_row, column=1).value is not None:
            data_row += 1

    ws.cell(row=data_row, column=1, value=as_of)
    for i, s in enumerate(result["score_list"], start=2):
        ws.cell(row=data_row, column=i, value=s)
    total_col = 2 + len(CATEGORIES)
    grade_col = total_col + 1
    notes_col = grade_col + 1
    from openpyxl.utils import get_column_letter
    first_letter = get_column_letter(2)
    last_letter = get_column_letter(1 + len(CATEGORIES))
    ws.cell(row=data_row, column=total_col,
            value=f"=SUM({first_letter}{data_row}:{last_letter}{data_row})")
    # Grade via an inline lookup table written once per call (self-contained,
    # doesn't depend on the Reference sheet's bands being present/identical).
    band_pairs = ",".join(f'{{{m},"{g}"}}' for m, g in reversed(GRADE_BANDS))
    # Simpler & robust across LibreOffice: nested IFs instead of a literal array.
    expr = '"Reject"'
    for min_score, g in GRADE_BANDS[:-1][::-1]:
        expr = f'IF({get_column_letter(total_col)}{data_row}>={min_score},"{g}",{expr})'
    ws.cell(row=data_row, column=grade_col, value=f"={expr}")
    ws.cell(row=data_row, column=notes_col, value=result["notes"])

    wb.save(tracker_path)
    return data_row


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--show-categories", action="store_true")
    sub = p.add_subparsers(dest="command")

    sc = sub.add_parser("score", help="Score the portfolio's construction and print the scorecard.")
    sc.add_argument("--scores", required=True, help=f"{len(CATEGORIES)} comma-separated numbers, fixed category order.")
    sc.add_argument("--notes", default="")

    ah = sub.add_parser("analyze-holdings", help="Compute concentration/diversification stats from a holdings CSV.")
    ah.add_argument("csv_path", help="CSV columns: ticker,value,theme")

    fl = sub.add_parser("flags", help="Manage the tracker's Red Flag Log.")
    flsub = fl.add_subparsers(dest="flags_command")
    flist = flsub.add_parser("list")
    flist.add_argument("--tracker", required=True)
    flist.add_argument("--ticker")
    flist.add_argument("--status-contains")
    fadd = flsub.add_parser("add")
    fadd.add_argument("--tracker", required=True)
    fadd.add_argument("--flag", required=True)
    fadd.add_argument("--ticker", required=True)
    fadd.add_argument("--status", required=True)
    fadd.add_argument("--notes", default="")
    fadd.add_argument("--source", default="")
    fres = flsub.add_parser("resolve")
    fres.add_argument("--tracker", required=True)
    fres.add_argument("--ticker", required=True)
    fres.add_argument("--match", required=True, help="Substring to match against the flag/notes text.")
    fres.add_argument("--status", required=True)
    fres.add_argument("--notes", default="")

    ap = sub.add_parser("append", help="Log a dated Portfolio Construction Score row into the tracker.")
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--scores", required=True)
    ap.add_argument("--notes", default="")
    ap.add_argument("--as-of", help="YYYY-MM-DD; defaults to today.")

    args = p.parse_args()

    if args.show_categories:
        for name, cap in CATEGORIES:
            print(f"  {name:<55} /{cap}")
        print(f"  {'TOTAL CAP':<55} /100")
        return

    if args.command == "score":
        try:
            result = score_portfolio(parse_scores_arg(args.scores), args.notes)
        except ScoreError as e:
            sys.exit(f"Score error: {e}")
        print(format_scorecard(result))

    elif args.command == "analyze-holdings":
        rows = []
        with open(args.csv_path, newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append({"ticker": row["ticker"], "value": float(row["value"]), "theme": row["theme"]})
        print(format_holdings_analysis(analyze_holdings(rows)))

    elif args.command == "flags":
        if args.flags_command == "list":
            rows = list_flags(args.tracker, args.ticker, args.status_contains)
            if not rows:
                print("No matching flags found.")
            for r in rows:
                print(f"[{r['Ticker']}] {r['Flag']} -- {r['Status']} (last checked {r['Last Checked']})")
                if r["Finding / Notes"]:
                    print(f"    {r['Finding / Notes']}")
        elif args.flags_command == "add":
            row = add_flag(args.tracker, args.flag, args.ticker, args.status, args.notes, args.source)
            print(f"Added flag at row {row} in {args.tracker}.")
        elif args.flags_command == "resolve":
            rows = resolve_flag(args.tracker, args.ticker, args.match, args.status, args.notes)
            print(f"Updated row(s) {rows} in {args.tracker}.")
        else:
            fl.print_help()

    elif args.command == "append":
        try:
            result = score_portfolio(parse_scores_arg(args.scores), args.notes)
        except ScoreError as e:
            sys.exit(f"Score error: {e}")
        row = append_score_history(args.tracker, result, args.as_of)
        print(format_scorecard(result))
        print(f"\nLogged to row {row} in {args.tracker}'s 'Portfolio Construction Score' history table.")
        print("Run the xlsx skill's recalc.py on the tracker next so Total/Grade get cached values.")

    else:
        p.print_help()

    print("\nGolden Rule: a score measures structural quality/risk. It is never a buy or sell signal.")


if __name__ == "__main__":
    main()
